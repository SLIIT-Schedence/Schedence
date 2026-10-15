"""Read the approved proportional PP1 layout without saving or evaluating Excel."""

from dataclasses import dataclass, field
from decimal import Decimal, ROUND_HALF_UP
from hashlib import sha256
from pathlib import Path
import warnings
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from schedence_teaching.excel_import.records import ImportIssue
from schedence_teaching.excel_import.staging import SourceLocation
from schedence_teaching.schemas import Stream

from .students import CohortCount, GroupSizingPolicy, StudentCountImport, generate_groups


SHEET = "STUDENT_ALLOCATIONS"


@dataclass
class SpecializationAllocations:
    cohorts: list[CohortCount] = field(default_factory=list)
    issues: list[ImportIssue] = field(default_factory=list)
    specialization_names: dict[str, str] = field(default_factory=dict)
    sizing: GroupSizingPolicy | None = None
    source_sha256: str | None = None
    allocation_rows: int = 0
    zero_allocations: int = 0
    reconciled_cohorts: int = 0
    recomputed_formula_cells: int = 0

    @property
    def valid(self):
        return bool(self.cohorts) and self.sizing is not None and not any(row.severity == "error" for row in self.issues)

    def summary(self):
        return dict(valid=self.valid, allocation_status="USER_APPROVED_PROPORTIONAL_PP1_ESTIMATE",
            official_specialization_enrollment=False, source_sha256=self.source_sha256,
            allocation_rows=self.allocation_rows, positive_allocations=len(self.cohorts), zero_allocations=self.zero_allocations,
            reconciled_cohorts=self.reconciled_cohorts, total_students=sum(row.student_count for row in self.cohorts),
            specialization_names=self.specialization_names, recomputed_formula_cells=self.recomputed_formula_cells,
            issues=[vars(row) for row in self.issues])

    def groups(self):
        if not self.valid:
            raise ValueError("Specialization allocations must reconcile before generating groups.")
        return generate_groups(self.cohorts, self.sizing)


def import_specialization_allocations(path: Path, original: StudentCountImport) -> SpecializationAllocations:
    """Validate literals and only the inspected cumulative-rounding formula layout.

    Formula caches are cross-checked against exact Decimal ROUND_HALF_UP math.
    Missing caches are permitted only for these recognized allocation formulas.
    This is not an arbitrary Excel formula interpreter; source bytes are unchanged.
    """
    path = Path(path)
    result = SpecializationAllocations()

    def issue(code, message, row=None, column=None):
        result.issues.append(ImportIssue("error", code, path.name, SHEET, row, column, message))

    def integer(value, row, column, minimum=0):
        if type(value) is not int or value < minimum:
            issue("invalid_allocation_integer", "Supply a literal integer in the permitted range.", row, column)
            return None
        return value

    if not original.valid or path.suffix.lower() != ".xlsx":
        issue("invalid_allocation_source", "Supply an .xlsx allocation workbook and valid original aggregate counts.")
        return result
    try:
        result.source_sha256 = sha256(path.read_bytes()).hexdigest()
        sheets = []
        with warnings.catch_warnings():
            # openpyxl warns about unsupported visual extensions even for read-only
            # access. No save occurs, so no extension or formatting is removed.
            warnings.filterwarnings("ignore", message=".*extension is not supported and will be removed", category=UserWarning)
            for cached in (False, True):
                workbook = load_workbook(path, read_only=True, data_only=cached, keep_links=False)
                try:
                    sheets.append(list(workbook[SHEET].iter_rows(values_only=True)))
                finally:
                    workbook.close()
        raw, cached = sheets

        def cell(rows, row, column):
            return rows[row-1][column-1] if row <= len(rows) and column <= len(rows[row-1]) else None

        def value(row, column):
            return cell(raw, row, column)

        required = {(7, 1): "Max main group", (7, 4): "Max lab subgroup", (7, 7): "Lecture hall planning cap",
            (9, 1): "Code", (9, 2): "Specialization", (9, 3): "Editable Weight",
            (20, 1): "Study Year", (20, 2): "Semester", (20, 3): "Stream", (20, 4): "Editable Student Total",
            **{(41, index): header for index, header in enumerate(("Cohort ID", "Year", "Sem", "Stream", "Specialization", "Weight", "Students (allocated)"), 1)}}
        for (row, column), expected in required.items():
            if value(row, column) != expected:
                issue("allocation_layout", "Required source header does not match the inspected allocation layout.", row, str(column))
        if result.issues:
            return result
        try:
            result.sizing = GroupSizingPolicy(main_group_maximum=value(7, 2), lab_subgroup_maximum=value(7, 5), combined_planning_maximum=value(7, 8))
        except ValueError:
            issue("invalid_allocation_caps", "Allocation group caps must be positive literal integers.", 7)
        codes, weights = [], []
        for row in range(10, 18):
            code, name, weight = value(row, 1), value(row, 2), value(row, 3)
            if not isinstance(code, str) or not code.strip() or code in codes or not isinstance(name, str) or not name.strip():
                issue("invalid_specialization", "Supply distinct nonblank specialization codes and names.", row, "Code")
                continue
            if type(weight) not in {int, float} or not Decimal(str(weight)).is_finite() or weight < 0:
                issue("invalid_allocation_weight", "Weights must be finite nonnegative literal numbers.", row, "Editable Weight")
                continue
            codes.append(code)
            weights.append(Decimal(str(weight)))
            result.specialization_names[code] = name
        if len(codes) != 8 or sum(weights) <= 0:
            issue("allocation_weight_total", "The eight source weights must have a positive total.")
            return result
        expected_scopes = {(row.academic_year, row.semester, row.stream): row.student_count for row in original.cohorts}
        seen_scopes, seen_ids = set(), set()
        for cohort_index, input_row in enumerate(range(21, 37)):
            year, semester, stream = value(input_row, 1), value(input_row, 2), value(input_row, 3)
            total = integer(value(input_row, 4), input_row, "Editable Student Total", 1)
            try:
                if type(year) is not int or type(semester) is not int:
                    raise ValueError
                scope = (year, semester, Stream(stream))
            except ValueError:
                issue("invalid_cohort_scope", "Study year, semester and stream must match the original cohort scope.", input_row)
                continue
            if scope in seen_scopes or expected_scopes.get(scope) != total:
                issue("cohort_total_mismatch", "Duplicate cohort or total differs from the original aggregate workbook.", input_row, "Editable Student Total")
                continue
            seen_scopes.add(scope)
            allocated, prefix, previous = 0, Decimal(0), 0
            for index, (code, weight) in enumerate(zip(codes, weights)):
                row = 42 + cohort_index * 8 + index
                result.allocation_rows += 1
                prefix += weight
                rounded = int((Decimal(total) * prefix / sum(weights)).quantize(Decimal(1), rounding=ROUND_HALF_UP))
                expected_count, previous = rounded - previous, rounded
                raw_count = value(row, 7)
                expected_formula = f"=ROUND($D${input_row}*SUM($C$10:$C${10+index})/SUM($C$10:$C$17),0)"
                if index:
                    expected_formula += f"-ROUND($D${input_row}*SUM($C$10:$C${9+index})/SUM($C$10:$C$17),0)"
                if isinstance(raw_count, str) and raw_count.startswith("="):
                    if raw_count != expected_formula or value(row, 6) != f"=$C${10+index}":
                        issue("unsupported_allocation_formula", "Formula differs from the inspected proportional layout; no arbitrary formula is evaluated.", row, "Students (allocated)")
                        continue
                    cached_count = cell(cached, row, 7)
                    if cached_count is not None and (type(cached_count) is not int or cached_count != expected_count):
                        issue("stale_allocation_cache", "Cached allocation differs from the saved literal weights/totals; recalculate the source in Excel.", row, "Students (allocated)")
                        continue
                    count = expected_count
                    result.recomputed_formula_cells += 1
                else:
                    count = integer(raw_count, row, "Students (allocated)")
                if count is None:
                    continue
                if (value(row, 2), value(row, 3), value(row, 4), value(row, 5)) != (year, semester, stream, code):
                    issue("allocation_scope_mismatch", "Allocation year/semester/stream/specialization differs from its source cohort and code order.", row)
                    continue
                source_id = value(row, 1)
                identity = (source_id, code)
                if not isinstance(source_id, str) or not source_id.strip() or identity in seen_ids:
                    issue("duplicate_allocation", "Allocation cohort IDs and specialization records must be distinct.", row, "Cohort ID")
                    continue
                seen_ids.add(identity)
                allocated += count
                if count:
                    result.cohorts.append(CohortCount(cohort_id=f"PP1:{source_id}:{code}", academic_year=year, semester=semester,
                        stream=stream, student_count=count, specialization=code, source=SourceLocation(path.name, SHEET, row),
                        source_label="User-approved proportional allocation for PP1 testing; not exact specialization enrollment"))
                else:
                    result.zero_allocations += 1
            if allocated != total:
                issue("allocation_total_mismatch", "Specialization allocations do not add to the exact original cohort total.", input_row)
            else:
                result.reconciled_cohorts += 1
        if seen_scopes != set(expected_scopes):
            issue("allocation_cohort_coverage", "Allocation workbook must cover every original cohort exactly once.")
        if sha256(path.read_bytes()).hexdigest() != result.source_sha256:
            issue("allocation_source_changed", "Source changed during reading; retry with a stable saved workbook.")
    except (OSError, ValueError, KeyError, EOFError, BadZipFile, InvalidFileException, ParseError):
        issue("unreadable_allocation_workbook", "Cannot read the inspected specialization allocation workbook; source was not modified.")
    return result
