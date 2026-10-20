"""Read literal cohort inputs; generate planned groups without population claims."""

from dataclasses import dataclass, field
from hashlib import sha256
from itertools import product
from pathlib import Path
from urllib.parse import quote
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from pydantic import ValidationError
from openpyxl.utils.exceptions import InvalidFileException

from schedence_teaching.excel_import.source_reader import read_source
from schedence_teaching.excel_import.records import ImportIssue
from schedence_teaching.excel_import.staging import SourceLocation
from schedence_teaching.excel_import.wide_config import DatasetSheet
from schedence_teaching.preparation import build_group_hierarchy
from schedence_teaching.schemas import Stream, StudentGroup
from schedence_teaching.schemas.models import AcademicYear, InputContract, PositiveInteger, RequiredText, Semester


class GroupSizingPolicy(InputContract):
    main_group_maximum: PositiveInteger = 120
    lab_subgroup_maximum: PositiveInteger = 60
    combined_planning_maximum: PositiveInteger = 240
    # Technical guard: fail, never truncate; not a university group-size rule.
    maximum_generated_records: PositiveInteger = 20_000


class CohortCount(InputContract):
    cohort_id: RequiredText
    academic_year: AcademicYear
    semester: Semester
    stream: Stream
    student_count: PositiveInteger
    specialization: RequiredText | None = None
    source: SourceLocation
    source_label: str | None = None


class StudentWorkbookLayout(InputContract):
    worksheet: RequiredText = "STUDENT_GROUPS"
    header_row: PositiveInteger = 9
    first_data_row: PositiveInteger = 10
    last_data_row: PositiveInteger = 25
    require_all_scopes: bool = True


@dataclass
class StudentCountImport:
    cohorts: list[CohortCount] = field(default_factory=list)
    issues: list[ImportIssue] = field(default_factory=list)
    sizing: GroupSizingPolicy | None = None
    source_sha256: str | None = None

    @property
    def valid(self):
        return self.sizing is not None and bool(self.cohorts) and not any(row.severity == "error" for row in self.issues)

    def summary(self):
        return dict(valid=self.valid, data_status="user_supplied_supervisor_accepted_development_counts",
            cohort_count=len(self.cohorts), total_students=sum(row.student_count for row in self.cohorts),
            missing_specialization_count=sum(row.specialization is None for row in self.cohorts),
            intake_mapping={"Regular Intake": 1, "June Intake": 2},
            sizing=self.sizing.model_dump() if self.sizing else None, source_sha256=self.source_sha256,
            issues=[vars(row) for row in self.issues], population_disjointness_verified=False)


COHORT_COLUMNS = dict(cohort_id="Cohort ID", academic_year="Study year", semester="Semester",
    stream="Stream", specialization="Specialization", student_count="Students", source_label="Data source")
INTAKE_SEMESTERS = {"Regular Intake": 1, "June Intake": 2}


def import_student_counts(path: Path, *, layout: StudentWorkbookLayout | None = None) -> StudentCountImport:
    """Use the inspected input table only, never cached formula group outputs.

    No fixed 6720 total or cohort counts enter production validation. Edited
    literal counts are product inputs; independent grouping enforces supplied caps.
    """
    path = Path(path)
    layout = StudentWorkbookLayout.model_validate((layout or StudentWorkbookLayout()).model_dump())
    result = StudentCountImport()

    def issue(code, message, row=None, column=None, severity="error"):
        result.issues.append(ImportIssue(severity, code, path.name, layout.worksheet, row, column, message))

    if path.suffix.lower() != ".xlsx":
        issue("unsupported_workbook", "Supply an .xlsx student-count workbook.")
        return result
    try:
        result.source_sha256 = sha256(path.read_bytes()).hexdigest()
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            if layout.worksheet not in workbook.sheetnames:
                issue("missing_worksheet", "Configured student input worksheet is missing.")
                return result
            sheet = workbook[layout.worksheet]
            if any(sheet[address].value != label for address, label in (
                    ("A4", "Main target"), ("C4", "Lab cap"), ("G4", "Large hall"))):
                issue("sizing_layout", "Sizing labels do not match the inspected editable workbook layout.", 4)
                return result
            cap_cells = {"main_group_maximum": "B4", "lab_subgroup_maximum": "D4", "combined_planning_maximum": "H4"}
            for field_name, address in cap_cells.items():
                if sheet[address].data_type in {"f", "e"}:
                    issue("literal_sizing_required", "Sizing inputs must be positive literal integers.", 4, address)
                    return result
            try:
                result.sizing = GroupSizingPolicy(**{key: sheet[address].value for key, address in cap_cells.items()})
            except ValidationError:
                issue("invalid_sizing", "Supply positive integer main/lab/planning caps.", 4)
        finally:
            workbook.close()
        source = DatasetSheet(workbook=path, worksheet=layout.worksheet, header_row=layout.header_row)
        # The generic reader rejects formulas only in mapped literal input fields.
        records = read_source(source, set(COHORT_COLUMNS.values()), result, Path("."))
        result.issues = [row for row in result.issues if not (
            row.code == "unsupported_cell" and row.row is not None and row.row > layout.last_data_row)]
        identities, scopes = set(), set()
        for record in records:
            if not layout.first_data_row <= record.location.row <= layout.last_data_row:
                # Formula-only output rows are explicitly outside the input table.
                continue
            if record.invalid_cells:
                continue
            payload = {key: record.source_values.get(header) for key, header in COHORT_COLUMNS.items()}
            if isinstance(payload["semester"], str):
                payload["semester"] = INTAKE_SEMESTERS.get(payload["semester"], payload["semester"])
            if payload["specialization"] == "":
                payload["specialization"] = None
            try:
                cohort = CohortCount(**payload, source=record.location)
            except ValidationError as error:
                for detail in error.errors(include_input=False, include_context=False):
                    issue("invalid_cohort", "Correct the literal cohort value using the existing year/semester/stream/count contract.",
                        record.location.row, COHORT_COLUMNS.get(detail["loc"][0]) if detail["loc"] else None)
                continue
            scope = (cohort.academic_year, cohort.semester, cohort.stream, cohort.specialization)
            if cohort.cohort_id in identities or scope in scopes:
                issue("duplicate_cohort", "Cohort IDs and scoped cohort records must be distinct; no counts are merged.", record.location.row, "Cohort ID")
            identities.add(cohort.cohort_id)
            scopes.add(scope)
            result.cohorts.append(cohort)
        if layout.require_all_scopes:
            expected = set(product(range(1, 5), (1, 2), Stream))
            actual = {(row.academic_year, row.semester, row.stream) for row in result.cohorts}
            if actual != expected:
                issue("cohort_scope_coverage", "Supply all study-year 1-4, semester 1-2 and weekday/weekend cohort scopes.")
        if any(row.specialization is None for row in result.cohorts):
            issue("specialization_not_supplied", "Unknown specialization remains absent; real offering/group conversion needs explicit mapping.", severity="warning")
        issue("aggregate_counts_not_population_proof", "Generated partitions are planning groups, not verified individual population disjointness.", severity="warning")
        if sha256(path.read_bytes()).hexdigest() != result.source_sha256:
            issue("source_changed", "Workbook changed during import; discard this snapshot and retry after the edit is saved.")
    except (OSError, ValueError, KeyError, EOFError, BadZipFile, InvalidFileException, ParseError):
        issue("unreadable_workbook", "Student input workbook could not be read; no source corrections were attempted.")
    return result


@dataclass(frozen=True)
class PlannedGroups:
    cohorts: tuple[CohortCount, ...]
    groups: tuple[StudentGroup, ...]
    cohort_by_group: dict[str, str]
    specializations: dict[str, str]
    sizing: GroupSizingPolicy

    def summary(self):
        main = [row for row in self.groups if row.parent_group_id is None]
        labs = [row for row in self.groups if row.parent_group_id is not None]
        return dict(total_students=sum(row.student_count for row in main), main_group_count=len(main),
            lab_subgroup_count=len(labs), group_record_count=len(self.groups),
            main_student_sum=sum(row.student_count for row in main), lab_student_sum=sum(row.student_count for row in labs),
            maximum_main_students=max((row.student_count for row in main), default=0),
            maximum_lab_students=max((row.student_count for row in labs), default=0),
            group_identity_status="deterministic_product_planning_ids_not_authoritative_database_ids",
            effective_sizing=self.sizing.model_dump(),
            population_disjointness_verified=False, specialization_inferred=False)


def balanced_sizes(total: int, maximum: int) -> tuple[int, ...]:
    if type(total) is not int or type(maximum) is not int or total <= 0 or maximum <= 0:
        raise ValueError("Student totals and maximum group sizes must be positive integers.")
    count = (total + maximum - 1) // maximum
    quotient, remainder = divmod(total, count)
    return tuple(quotient + (index < remainder) for index in range(count))


def generate_groups(cohorts: list[CohortCount] | tuple[CohortCount, ...], sizing: GroupSizingPolicy) -> PlannedGroups:
    cohorts = tuple(CohortCount.model_validate(row.model_dump()) for row in cohorts)
    sizing = GroupSizingPolicy.model_validate(sizing.model_dump())
    groups, links, specializations = [], {}, {}
    if len({row.cohort_id for row in cohorts}) != len(cohorts):
        raise ValueError("Supply distinct cohort IDs; no source cohort is duplicated or merged.")
    for cohort in cohorts:
        count = (cohort.student_count + sizing.main_group_maximum - 1) // sizing.main_group_maximum
        minimum_records = 2 * count
        if len(groups) + minimum_records > sizing.maximum_generated_records:
            raise ValueError("Technical group-record limit exceeded; no groups were truncated.")
        for index, size in enumerate(balanced_sizes(cohort.student_count, sizing.main_group_maximum), 1):
            parent = f"PLAN-Y{cohort.academic_year}-S{cohort.semester}-{cohort.stream.value}-{quote(cohort.cohort_id, safe='')}-G{index:03}"
            common = dict(academic_year=cohort.academic_year, semester=cohort.semester, stream=cohort.stream)
            group = StudentGroup(group_id=parent, student_count=size, **common)
            child_count = (size + sizing.lab_subgroup_maximum - 1) // sizing.lab_subgroup_maximum
            if len(groups) + 1 + child_count > sizing.maximum_generated_records:
                raise ValueError("Technical group-record limit exceeded; no groups were truncated.")
            children = [StudentGroup(group_id=f"{parent}.L{child_index:02}", student_count=child_size,
                parent_group_id=parent, **common) for child_index, child_size in
                enumerate(balanced_sizes(size, sizing.lab_subgroup_maximum), 1)]
            groups.extend([group, *children])
            for row in [group, *children]:
                links[row.group_id] = cohort.cohort_id
                if cohort.specialization is not None:
                    specializations[row.group_id] = cohort.specialization
    build_group_hierarchy(groups)
    return PlannedGroups(cohorts, tuple(groups), links, specializations, sizing)
