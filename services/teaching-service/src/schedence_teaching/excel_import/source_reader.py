"""Read source cells once for wide teaching and venue adapters."""

from pathlib import Path
from typing import Protocol
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .records import ImportIssue
from .staging import SourceLocation, SourceRecord
from .wide_config import DatasetSheet


class SourceIssueSink(Protocol):
    issues: list[ImportIssue]


def read_source(
    source: DatasetSheet, headers: set[str], result: SourceIssueSink, base_dir: Path,
) -> list[SourceRecord]:
    records = []
    workbook = None

    def issue(code, column, message):
        result.issues.append(ImportIssue(
            "error", code, source.workbook.name, source.worksheet, source.header_row, column, message,
        ))

    if source.workbook.suffix.lower() != ".xlsx":
        issue("unsupported_workbook", None, "Only .xlsx workbooks are supported.")
        return records
    try:
        path = source.workbook if source.workbook.is_absolute() else base_dir / source.workbook
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        if source.worksheet not in workbook.sheetnames:
            issue("missing_worksheet", None, "Configured worksheet does not exist.")
            return records
        sheet = workbook[source.worksheet]
        cells = next(sheet.iter_rows(min_row=source.header_row, max_row=source.header_row), ())
        actual_headers = [cell.value for cell in cells]
        bad_header = False
        for header in sorted(headers):
            if actual_headers.count(header) != 1:
                code = "missing_column" if header not in actual_headers else "duplicate_column"
                issue(code, header, "Required source header must occur exactly once.")
                bad_header = True
        if bad_header:
            return records
        for row_number, cells in enumerate(sheet.iter_rows(min_row=source.header_row + 1), source.header_row + 1):
            if all(cell.value is None for cell in cells):
                continue
            location = SourceLocation(source.workbook.name, source.worksheet, row_number)
            values = {header: cell.value for header, cell in zip(actual_headers, cells) if isinstance(header, str)}
            invalid = {header for header, cell in zip(actual_headers, cells)
                       if header in headers and cell.data_type in {"f", "e"}}
            for header in sorted(invalid):
                location.issue(result, "unsupported_cell", header,
                               "Mapped input must be a literal value, not a formula or Excel error.")
            records.append(SourceRecord(location, values, invalid))
    except (OSError, InvalidFileException, BadZipFile, ParseError, ValueError, KeyError, EOFError):
        issue("unreadable_workbook", None, "Source could not be read as an .xlsx workbook.")
        records = []
    finally:
        if workbook is not None:
            workbook.close()
    return records

