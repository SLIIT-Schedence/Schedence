"""Read mapped .xlsx cells without saving or evaluating source workbooks."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import StrictBool, StrictFloat, StrictInt, TypeAdapter, ValidationError

from schedence_teaching.schemas.models import InputContract, RequiredText

from .config import RECORD_MODELS, TableMapping
from .records import ImportIssue, ImportResult, LecturerPreferenceRecord

IDENTITY_FIELDS = {
    "staff": ("staff_id",),
    "preferences": ("staff_id",),
    "offerings": ("offering_id",),
    "teaching_teams": ("offering_id", "staff_id", "role"),
}
PREFERENCE_ADAPTERS = {
    "text": TypeAdapter(RequiredText),
    "integer": TypeAdapter(StrictInt),
    "number": TypeAdapter(StrictFloat, config={"allow_inf_nan": False}),
    "boolean": TypeAdapter(StrictBool),
}


def _blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


@dataclass
class SourceRow:
    table: TableMapping
    row: int
    record: InputContract | None
    identity: tuple[str, ...] | None

    def issue(
        self, result: ImportResult, code: str, field_name: str | None, message: str,
        *, severity: str = "error",
    ) -> None:
        column = self.table.columns.get(field_name)
        if field_name in self.table.preference_fields:
            column = self.table.preference_fields[field_name].column
        result.issues.append(ImportIssue(
            severity, code, self.table.workbook.name, self.table.worksheet,
            self.row, column, message,
        ))


def _parse_row(table: TableMapping, row_number: int, cells: tuple, positions: dict, result: ImportResult) -> SourceRow:
    payload = {}
    values = {}
    source = SourceRow(table, row_number, None, None)
    invalid_cell = False
    mappings = dict(table.columns)
    mappings.update({name: field.column for name, field in table.preference_fields.items()})
    for field_name, header in mappings.items():
        cell = cells[positions[header]]
        raw_value = cell.value
        if cell.data_type in {"f", "e"}:
            source.issue(result, "unsupported_cell", field_name,
                         "Formula and Excel error cells are not accepted; supply verified literal values.")
            invalid_cell = True
            continue
        if field_name in table.preference_fields:
            rule = table.preference_fields[field_name]
            if _blank(raw_value):
                if rule.required:
                    source.issue(result, "missing_value", field_name, "Required preference value is missing.")
                    invalid_cell = True
                continue
            try:
                PREFERENCE_ADAPTERS[rule.value_type].validate_python(raw_value)
            except ValidationError:
                source.issue(result, "invalid_preference_type", field_name,
                             f"Preference value must have the configured {rule.value_type} type.")
                invalid_cell = True
                continue
            values[field_name] = raw_value
        else:
            translations = table.value_maps.get(field_name, {})
            payload[field_name] = translations.get(raw_value, raw_value) if isinstance(raw_value, str) else raw_value

    key_values = [payload.get(field_name) for field_name in IDENTITY_FIELDS[table.kind]]
    if all(isinstance(value, str) and not _blank(value) for value in key_values):
        source.identity = tuple(key_values)
    if invalid_cell:
        return source

    if table.kind == "preferences":
        model = LecturerPreferenceRecord
        payload["values"] = values
    else:
        model = RECORD_MODELS[table.kind]
    try:
        source.record = model.model_validate(payload)
    except ValidationError as error:
        for detail in error.errors(include_input=False, include_url=False):
            location = detail["loc"]
            field_name = str(location[-1]) if location else None
            if field_name == "values":
                field_name = next(iter(table.preference_fields))
            message = detail["msg"]
            if field_name in payload and _blank(payload[field_name]):
                message = "Required value is missing."
            source.issue(result, "invalid_record", field_name, message)
    return source


def read_table(table: TableMapping, base_dir: Path, result: ImportResult) -> list[SourceRow]:
    """Read one explicitly mapped table, collecting cell-level validation issues."""
    rows: list[SourceRow] = []

    def issue(code: str, message: str, row=None, column=None, severity="error") -> None:
        result.issues.append(ImportIssue(
            severity, code, table.workbook.name, table.worksheet, row, column, message,
        ))

    if table.workbook.suffix.lower() != ".xlsx":
        issue("unsupported_workbook", "Only .xlsx workbooks are supported.")
        return rows
    path = table.workbook if table.workbook.is_absolute() else base_dir / table.workbook
    workbook = None
    try:
        # Keep formulas visible so they cannot masquerade as cached literal values.
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        if table.worksheet not in workbook.sheetnames:
            issue("missing_worksheet", "Configured worksheet does not exist.")
            return rows
        worksheet = workbook[table.worksheet]
        header_cells = next(worksheet.iter_rows(min_row=table.header_row, max_row=table.header_row), ())
        headers = [cell.value for cell in header_cells]
        mapped_headers = list(table.columns.values()) + [rule.column for rule in table.preference_fields.values()]
        bad_header = False
        for header in mapped_headers:
            occurrences = headers.count(header)
            if occurrences != 1:
                code = "missing_column" if occurrences == 0 else "duplicate_column"
                message = "Configured column is missing." if occurrences == 0 else "Mapped header occurs more than once."
                issue(code, message, table.header_row, header)
                bad_header = True
        if bad_header:
            return rows
        positions = {header: headers.index(header) for header in mapped_headers}
        if any(header is not None and header not in mapped_headers for header in headers):
            issue("unmapped_columns", "Unmapped columns are ignored; configure any fields that must be imported.",
                  table.header_row, severity="warning")
        for row_number, cells in enumerate(
            worksheet.iter_rows(min_row=table.header_row + 1, max_col=len(headers)),
            start=table.header_row + 1,
        ):
            if all(cell.value is None for cell in cells):
                continue
            result.rows_read[table.kind] += 1
            rows.append(_parse_row(table, row_number, cells, positions, result))
        if not rows:
            issue("empty_table", "Mapped table contains no data rows.", table.header_row, severity="warning")
    except (OSError, InvalidFileException, BadZipFile, ParseError, ValueError, KeyError, EOFError):
        issue("unreadable_workbook", "Workbook could not be read; check the path and that it is a valid .xlsx file.")
        # Never accept a partial table after an I/O or XML parsing failure.
        rows = []
    finally:
        if workbook is not None:
            workbook.close()
    return rows
