"""Validate the supplied review against source rows without reclassifying modules."""

from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

HEADERS = ("Source_Row", "Module_Code", "Module_Name", "Year", "Semester", "Specializations",
    "Lecture_h", "Tutorial_h", "Workshop_h", "Lab_h", "Total_Contact_h",
    "Elective_Title_Evidence", "Academic_Classification", "PP1_TEST_CANDIDATE", "Notes")
DIRECT = {"Module_Code": "Module_Code", "Module_Name": "Module_Name", "Year": "Year",
    "Semester": "Semester", "Specializations": "Specialization", "Lecture_h": "Lecture_Hours_Week",
    "Tutorial_h": "Tutorial_Hours_Week", "Workshop_h": "Workshop_Hours_Week", "Lab_h": "Lab_Hours_Week"}


@dataclass(frozen=True)
class ClassificationReview:
    source_row: int
    review_row: int
    academic_classification: str
    pp1_test_candidate: bool
    source_values: dict


def validate_review_rows(rows, source_modules):
    """Fail closed on coverage, duplicate rows or changes to original academic facts."""
    originals = {row.location.row: row for row in source_modules
        if row.contract_fields["academic_year"] == 4}
    result, seen = [], set()
    for worksheet_row, values in rows:
        if not values.get("Source_Row"):
            continue
        reference = values["Source_Row"]
        if type(reference) is not int or reference not in originals or reference in seen:
            raise ValueError(f"Year4_Classification row {worksheet_row}: unknown or repeated Source_Row")
        seen.add(reference)
        original = originals[reference].source_values
        for target, column in DIRECT.items():
            if values.get(target) != original.get(column):
                raise ValueError(f"Year4_Classification row {worksheet_row}, {target}: differs from module source")
        total = sum(Decimal(str(values[key])) for key in ("Lecture_h", "Tutorial_h", "Workshop_h", "Lab_h"))
        if total != Decimal(str(values.get("Total_Contact_h"))):
            raise ValueError(f"Year4_Classification row {worksheet_row}: inconsistent contact total")
        classification = values.get("Academic_Classification")
        candidate = values.get("PP1_TEST_CANDIDATE")
        if classification not in {"UNCLASSIFIED", "COMPULSORY", "ELECTIVE"} or candidate not in {"YES", "NO"}:
            raise ValueError(f"Year4_Classification row {worksheet_row}: invalid classification or candidate flag")
        result.append(ClassificationReview(reference, worksheet_row, classification, candidate == "YES", values))
    if seen != set(originals):
        raise ValueError("Year4_Classification does not cover every Year 4 source module exactly once")
    return tuple(result)


def read_classification_review(path: Path, source_modules):
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        sheet = workbook["Year4_Classification"]
        rows = list(sheet.iter_rows(values_only=True))
        if tuple(rows[4]) != HEADERS:
            raise ValueError("Year4_Classification row 5: unexpected headers")
        return validate_review_rows(((index, dict(zip(HEADERS, values)))
            for index, values in enumerate(rows[5:], 6)), source_modules)
    finally:
        workbook.close()
