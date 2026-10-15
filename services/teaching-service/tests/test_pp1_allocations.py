"""Sanitized allocation layouts and coordinated PP1 scheduling/failure exports."""

from copy import deepcopy
from dataclasses import replace
from decimal import Decimal
from hashlib import sha256
import json
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
import pytest

from schedence_teaching.excel_import import VenueImportResult
from schedence_teaching.excel_import.staging import SourceLocation, SourceModuleRecord, SourceRecord, SourceTeamAssignment, WideImportResult
from schedence_teaching.integration import assess_readiness
from schedence_teaching.pilot.allocations import SpecializationAllocations, import_specialization_allocations
from schedence_teaching.pilot.pp1 import PP1Configuration, attendance_capacity_diagnostics, prepare_pp1
from schedence_teaching.pilot.pp1_run import export_pp1, prove_attendance_infeasibility, run_pp1
from schedence_teaching.pilot.students import CohortCount, GroupSizingPolicy, import_student_counts
from schedence_teaching.schemas import StaffMember, TeachingRole, Venue
from schedence_teaching.solver import SolverInputError, solve_timetable
from schedence_teaching.solver.evaluation import validate_evaluation_result

from test_authoritative_integration import complete_template, INSTITUTION
from test_student_group_pilot import workbook_path


CODES = ("IT", "SE", "DS", "CYBER", "ISE", "CSNE", "COMPUTER_SCIENCE", "IM")
WEIGHTS = (450, 300, 300, 200, 200, 100, 100, 50)


def write_allocations(file, original, weights=WEIGHTS):
    book = Workbook()
    sheet = book.active
    sheet.title = "STUDENT_ALLOCATIONS"
    for address, value in {"A7": "Max main group", "B7": 120, "D7": "Max lab subgroup", "E7": 60,
        "G7": "Lecture hall planning cap", "H7": 240, "A9": "Code", "B9": "Specialization", "C9": "Editable Weight"}.items():
        sheet[address] = value
    for index, (code, weight) in enumerate(zip(CODES, weights), 10):
        sheet.cell(index, 1, code)
        sheet.cell(index, 2, f"Synthetic {code}")
        sheet.cell(index, 3, weight)
    for column, value in enumerate(("Study Year", "Semester", "Stream", "Editable Student Total"), 1):
        sheet.cell(20, column, value)
    for column, value in enumerate(("Cohort ID", "Year", "Sem", "Stream", "Specialization", "Weight", "Students (allocated)"), 1):
        sheet.cell(41, column, value)
    for index, cohort in enumerate(original.cohorts):
        input_row = 21 + index
        for column, value in enumerate((cohort.academic_year, cohort.semester, cohort.stream.value, cohort.student_count), 1):
            sheet.cell(input_row, column, value)
        for position, code in enumerate(CODES):
            row = 42 + index * 8 + position
            formula = f"=ROUND($D${input_row}*SUM($C$10:$C${10+position})/SUM($C$10:$C$17),0)"
            if position:
                formula += f"-ROUND($D${input_row}*SUM($C$10:$C${9+position})/SUM($C$10:$C$17),0)"
            for column, value in enumerate((cohort.cohort_id, cohort.academic_year, cohort.semester, cohort.stream.value,
                code, f"=$C${10+position}", formula), 1):
                sheet.cell(row, column, value)
    sheet["A175"] = "Generated group details are not used as academic inputs"
    book.save(file)
    book.close()


@pytest.fixture
def allocation_source(workbook_path, tmp_path):
    original = import_student_counts(workbook_path)
    file = tmp_path / "synthetic-allocation.xlsx"
    write_allocations(file, original)
    return file, original


def edit(file, address, value):
    book = load_workbook(file)
    book.active[address] = value
    book.save(file)
    book.close()


def test_allocations_reconcile_all_scopes_and_automatically_map_groups(allocation_source):
    file, original = allocation_source
    before = sha256(file.read_bytes()).hexdigest()
    result = import_specialization_allocations(file, original)
    assert result.valid, result.summary()
    assert result.reconciled_cohorts == 16 and result.allocation_rows == 128
    assert result.recomputed_formula_cells == 128  # No Excel calculation runtime/caches required.
    assert [row.student_count for row in result.cohorts[:8]] == [191, 127, 127, 84, 85, 42, 43, 21]
    plan = result.groups()
    assert plan.summary()["total_students"] == 6720
    assert (plan.summary()["main_group_count"], plan.summary()["lab_subgroup_count"]) == (140, 198)
    assert set(plan.specializations.values()) == set(CODES)
    assert all(row.student_count <= (120 if row.parent_group_id is None else 60) for row in plan.groups)
    assert not result.summary()["official_specialization_enrollment"]
    assert before == result.source_sha256 == sha256(file.read_bytes()).hexdigest()


def test_excel_half_up_rounding_is_preserved_not_python_bankers_rounding(allocation_source):
    file, original = allocation_source
    original.cohorts[0].student_count = 4
    write_allocations(file, original, weights=(1,)*8)
    result = import_specialization_allocations(file, original)
    assert result.valid
    assert result.zero_allocations == 4
    assert [(row.specialization, row.student_count) for row in result.cohorts[:4]] == [("IT", 1), ("DS", 1), ("ISE", 1), ("COMPUTER_SCIENCE", 1)]


@pytest.mark.parametrize("address,value,code", [("D21", 721, "cohort_total_mismatch"), ("G42", -1, "invalid_allocation_integer"),
    ("G42", 1, "allocation_total_mismatch"), ("C10", -1, "invalid_allocation_weight"), ("C10", "450", "invalid_allocation_weight"),
    ("C10", "=450", "invalid_allocation_weight"), ("A10", "SE", "invalid_specialization"),
    ("B42", 2, "allocation_scope_mismatch"), ("D42", "WEEKEND", "allocation_scope_mismatch"),
    ("E42", "UNKNOWN", "allocation_scope_mismatch"), ("G42", "=SUM(A1:A4)", "unsupported_allocation_formula"),
    ("F42", "=$C$11", "unsupported_allocation_formula"), ("B7", 0, "invalid_allocation_caps"), ("A41", "Unknown", "allocation_layout")])
def test_invalid_allocations_report_exact_source_location(allocation_source, address, value, code):
    file, original = allocation_source
    edit(file, address, value)
    result = import_specialization_allocations(file, original)
    assert not result.valid
    assert any(issue.code == code and issue.workbook == file.name and issue.worksheet == "STUDENT_ALLOCATIONS" for issue in result.issues)
    with pytest.raises(ValueError):
        result.groups()


def test_literal_allocation_edits_preserved_without_inventing_zero_groups(allocation_source):
    file, original = allocation_source
    book = load_workbook(file)
    for index in range(8):
        book.active.cell(42+index, 7, 720 if index == 0 else 0)
    book.save(file)
    book.close()
    result = import_specialization_allocations(file, original)
    assert result.valid
    assert result.zero_allocations == 7
    assert result.cohorts[0].student_count == 720 and result.cohorts[0].specialization == "IT"


def test_stale_cached_formula_values_are_rejected(allocation_source):
    file, original = allocation_source
    buffer = BytesIO()
    with ZipFile(file) as source, ZipFile(buffer, "w") as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(content)
                namespace = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                node = root.find(".//m:c[@r='G42']/m:v", namespace)
                node.text = "999"
                content = ET.tostring(root)
            target.writestr(entry, content)
    file.write_bytes(buffer.getvalue())
    result = import_specialization_allocations(file, original)
    assert not result.valid and any(row.code == "stale_allocation_cache" for row in result.issues)


def test_broken_allocation_file_never_looks_like_a_zero_record_success(allocation_source):
    file, original = allocation_source
    file.write_text("not an xlsx", encoding="utf-8")
    result = import_specialization_allocations(file, original)
    assert not result.valid and not result.cohorts
    assert any(row.code == "unreadable_allocation_workbook" for row in result.issues)


def configuration():
    file = Path(__file__).resolve().parents[1] / "examples" / "pp1-test-calendar.json"
    raw = json.loads(file.read_text(encoding="utf-8"))
    raw["solver_config"]["max_time_seconds"] = 5.0
    return PP1Configuration.model_validate(raw)


def small_batch(*, position="Lecturer", lab=0, lecture=2, config=None, allocation_specializations=None):
    source = WideImportResult(source_confirmation_reference="synthetic:confirmed-source", lecture_hours_interpretation="ONE_CONTINUOUS_WEEKLY_LECTURE")
    source.staff_members = [StaffMember(staff_id="SYN-LECTURER", staff_name="Synthetic Name", position=position, staff_type="ACADEMIC")]
    allocations = SpecializationAllocations(sizing=GroupSizingPolicy(), specialization_names={"IT": "Synthetic IT", "SE": "Synthetic SE"})
    for index, code in enumerate(("IT", "SE"), 2):
        location = SourceLocation("synthetic-modules.xlsx", "Modules_183", index)
        source.source_modules.append(SourceModuleRecord(location, dict(Lecture_Hours_Week=lecture, Tutorial_Hours_Week=0, Workshop_Hours_Week=0, Lab_Hours_Week=lab),
            contract_fields=dict(module_code="IT2100", module_name=f"Synthetic {code}", academic_year=1, semester=1, specialization=code, degree_type="Synthetic")))
        source.source_assignments.append(SourceTeamAssignment(location, "LIC_Staff_ID", "LIC", "SYN-LECTURER", "Synthetic", TeachingRole.LIC))
        allocations.cohorts.append(CohortCount(cohort_id=f"SYN-{code}", academic_year=1, semester=1, stream="WEEKDAY", student_count=80,
            specialization=allocation_specializations[index-2] if allocation_specializations else code,
            source=SourceLocation("synthetic-allocation.xlsx", "STUDENT_ALLOCATIONS", 40+index)))
    venues = VenueImportResult()
    for identity, kind in (("SYN-HALL", "Lecture Hall"), ("SYN-LAB", "Computer Lab")):
        venues.venues.append(Venue(venue_id=identity, venue_name="Synthetic Room", venue_type=kind, capacity=120))
        venues.source_venues.append(SourceRecord(SourceLocation("synthetic-venues.xlsx", "Venue_Master", len(venues.venues)+1), {}))
    return prepare_pp1(allocations, source, venues, config or configuration())


def test_pp1_preserves_all_source_teams_and_does_not_enable_combining_or_fake_ids():
    batch = small_batch(lab=2)
    teaching = batch.prepared.input_snapshot
    assert batch.metadata["shared_lecture_combining"] is False
    assert batch.metadata["cross_specialization_overlap_optimization"] is False
    assert len(teaching.teaching_team_members) == 2
    assert all(row.staff_id == "SYN-LECTURER" for row in teaching.teaching_team_members)
    assert all(row.offering_id.startswith("PP1:source-row%3A") for row in teaching.module_offerings)
    assert all(row.duration_minutes == 120 and row.sessions_per_week == 1 for row in teaching.session_requirements)
    assert sum(row.session_type.value == "LECTURE" for row in teaching.session_requirements) == 2
    assert sum(row.session_type.value == "LAB" for row in teaching.session_requirements) == 4


def test_edited_unknown_specialization_cannot_silently_lose_module_demand():
    with pytest.raises(ValueError, match="no source module mapping"):
        small_batch(allocation_specializations=("IT", "UNKNOWN-CODE"))


def test_pp1_lunch_scope_and_configured_venues_preserve_both_streams():
    batch = small_batch()
    assert batch.configuration.calendar.week_start.isoformat() == "2026-10-19"
    lunch = batch.solver_inputs["student_lunch_policy"]
    assert {row.value for row in lunch.applicable_streams} == {"WEEKDAY", "WEEKEND"}
    assert len(lunch.applicable_dates) == 7
    assert [(row.window.start_minute, row.window.end_minute) for row in lunch.year_periods] == [(750, 810), (750, 810), (810, 870), (810, 870)]
    assert all(len(row.available_windows) == 7 and row.blocked_periods == [] for row in batch.solver_inputs["venue_input"].calendars)


@pytest.fixture(scope="module")
def solved_small(tmp_path_factory):
    batch = small_batch()
    result, report = run_pp1(batch, tmp_path_factory.mktemp("pp1") / "valid")
    assert report["complete_timetable_generated"]
    return batch, result, report


def test_one_native_solve_prevents_cross_specialization_lecturer_and_venue_clashes(solved_small):
    batch, result, report = solved_small
    validate_evaluation_result(batch.prepared, batch.config, result, **batch.solver_inputs)
    assert report["scheduled_sessions"] == 2 and report["unscheduled_sessions"] == 0
    assert report["clash_validation"] == "PASSED"
    assert not result.combination_enabled
    first, second = result.schedule
    assert first.day != second.day or first.end_minute <= second.start_minute or second.end_minute <= first.start_minute


def test_export_rejects_corruption_before_creating_output_files(solved_small, tmp_path):
    batch, result, _ = solved_small
    with pytest.raises(ValueError):
        export_pp1(batch, replace(result, schedule=result.schedule[:-1]), tmp_path / "bad")
    assert not (tmp_path / "bad").exists()


def test_unclassified_testing_option_preserves_staff_without_inventing_rank(tmp_path):
    batch = small_batch(position="Head-CSE")
    assert batch.solver_inputs["staff_input"].allow_unclassified_staff_for_testing is True
    assert not batch.solver_inputs["staff_input"].workload_input.profiles
    result, report = run_pp1(batch, tmp_path / "head")
    assert report["complete_timetable_generated"]
    assert result.workload_report.unassessed_staff_ids == ("SYN-LECTURER",)
    assert batch.prepared.input_snapshot.staff_members[0].position == "Head-CSE"
    strict = deepcopy(batch.solver_inputs)
    strict["staff_input"].allow_unclassified_staff_for_testing = False
    with pytest.raises(SolverInputError, match="workload profiles"):
        solve_timetable(batch.prepared, batch.config, **strict)


def test_actual_campus_readiness_rejects_the_testing_workload_option(complete_template):
    submission = deepcopy(complete_template)
    submission.inputs["mode"] = "ACTUAL_CAMPUS"
    submission.inputs["calendar"]["synthetic_data"] = False
    submission.inputs["staff_input"]["allow_unclassified_staff_for_testing"] = True
    assessed = assess_readiness(submission, expected_institution_id=INSTITUTION, actual_campus=True)
    assert not assessed.report.ready_for_actual_campus
    assert any(row.code == "testing_workload_override" for row in assessed.report.diagnostics)


def test_existing_supplied_availability_and_caps_are_not_replaced():
    batch = small_batch()
    config = batch.configuration.model_copy(deep=True)
    config.staff_input = batch.solver_inputs["staff_input"].model_copy(deep=True)
    config.staff_input.availability[0].blocked_periods = []
    config.staff_input.availability[0].weekly_teaching_max_minutes = 0
    # The native caller still enforces this independently of workload targets.
    contexts = {**batch.solver_inputs, "staff_input": config.staff_input}
    result = solve_timetable(batch.prepared, batch.config, **contexts)
    assert result.status.value == "INFEASIBLE"
    assert not result.schedule


def test_supplied_blocked_period_remains_a_hard_staff_restriction():
    from schedence_teaching.solver.staff_inputs import StaffBlockedInterval
    batch = small_batch()
    staff = deepcopy(batch.solver_inputs["staff_input"])
    staff.availability[0].blocked_periods = [StaffBlockedInterval(day="MONDAY", start_minute=510, end_minute=1050)]
    contexts = {**batch.solver_inputs, "staff_input": staff}
    result = solve_timetable(batch.prepared, batch.config, **contexts)
    validate_evaluation_result(batch.prepared, batch.config, result, **contexts)
    assert result.schedule and all(row.day != "MONDAY" for row in result.schedule)


def test_time_capacity_proof_includes_parent_lectures_and_each_child_lab():
    batch = small_batch(lab=2)
    # A single 3-hour opening day leaves 2 teaching hours after lunch.
    config = batch.configuration.model_dump()
    config["calendar"]["day_dates"] = {"MONDAY": "2026-10-19"}
    config["solver_config"]["days"] = [dict(day="MONDAY", windows={"WEEKDAY": [dict(start_minute=690, end_minute=870)]})]
    changed = PP1Configuration.model_validate(config)
    limited = replace(batch, config=changed.solver_config, configuration=changed)
    diagnostics = attendance_capacity_diagnostics(limited)
    assert len(diagnostics) == 4
    assert all(row["required_minutes"] == 240 and row["available_minutes"] == 120 for row in diagnostics)


def test_failed_run_exports_every_unscheduled_occurrence_and_no_fake_assignments(tmp_path):
    batch = small_batch()
    staff = deepcopy(batch.solver_inputs["staff_input"])
    staff.availability[0].weekly_teaching_max_minutes = 0
    batch = replace(batch, solver_inputs={**batch.solver_inputs, "staff_input": staff})
    result, report = run_pp1(batch, tmp_path / "failed")
    assert result.status.value == "INFEASIBLE"
    assert report["scheduled_sessions"] == 0 and report["unscheduled_sessions"] == len(batch.prepared.sessions)
    assert not report["independently_validated"]
    assert report["failure_output_integrity_validated"]
    assert report["clash_validation"] == "NOT_APPLICABLE_NO_TIMETABLE"
    import csv
    with (tmp_path / "failed" / "unscheduled.csv").open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert all(not row["lecturer_id"] and not row["venue_id"] and not row["start_time"] for row in rows)
    assert len(rows) == 2
    views = json.loads((tmp_path / "failed" / "views.json").read_text(encoding="utf-8"))
    assert views["single_global_run"]
    assert set(views["views"]["specialization"]) == {"IT", "SE"}
    assert (tmp_path / "failed" / "specializations" / "IT.html").exists()
    assert "No complete timetable" in (tmp_path / "failed" / "timetable.html").read_text(encoding="utf-8")


@pytest.mark.parametrize("change", ["real-calendar", "missing-contact-type", "missing-date", "protected-without-config"])
def test_incomplete_configuration_is_rejected(change):
    raw = configuration().model_dump()
    if change == "real-calendar":raw["calendar"]["synthetic_data"] = False
    if change == "missing-contact-type":raw["session_audiences"].pop("LECTURE")
    if change == "missing-date":raw["calendar"]["day_dates"].pop("MONDAY")
    if change == "protected-without-config":raw["protected_period_config"] = {}
    with pytest.raises(ValueError):
        PP1Configuration.model_validate(raw)


def test_native_infeasible_source_subset_is_diagnostic_not_a_separate_schedule(tmp_path):
    raw = configuration().model_dump()
    raw["calendar"]["day_dates"] = {"MONDAY": "2026-10-19"}
    raw["solver_config"]["days"] = [dict(day="MONDAY", windows={"WEEKDAY": [dict(start_minute=690, end_minute=870)]})]
    batch = small_batch(lab=2, config=PP1Configuration.model_validate(raw))
    before = batch.prepared.input_snapshot.model_dump()
    result, report = run_pp1(batch, tmp_path / "certificate")
    assert result.status.value == "INFEASIBLE"
    assert report["independent_attendance_capacity_status"] == "INFEASIBLE"
    assert report["diagnostic_witness_status"] == "INFEASIBLE"
    witness = json.loads((tmp_path / "certificate" / "infeasibility-witness.json").read_text(encoding="utf-8"))
    assert witness["purpose"] == "INFEASIBILITY_DIAGNOSTIC_ONLY_NOT_A_SEPARATE_TIMETABLE"
    assert witness["source_subset_validated"] and witness["fixed_teaching_windows_unchanged"]
    assert witness["source_occurrences"] == 2
    assert set(witness["source_requirement_ids"]) <= {row.requirement_id for row in batch.prepared.input_snapshot.session_requirements}
    assert batch.prepared.input_snapshot.model_dump() == before


def test_capacity_diagnostic_does_not_run_a_witness_when_no_violation_exists():
    assert prove_attendance_infeasibility(small_batch()) is None
