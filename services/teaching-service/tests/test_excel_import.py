"""All workbook headers and records here are synthetic, not real dataset mappings."""

from __future__ import annotations

import json
import subprocess
import sys
from copy import deepcopy

import pytest
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from schedence_teaching.excel_import import ImportConfig, PreferenceField, TableMapping, import_excel
from schedence_teaching.excel_import.cli import main
from schedence_teaching.schemas import ModuleOffering, StaffMember, TeachingRole


@pytest.fixture
def excel_dataset(tmp_path, schema_payloads):
    path = tmp_path / "synthetic.xlsx"
    staff_columns = {
        "staff_id": "Synthetic ID", "staff_name": "Synthetic name",
        "position": "Synthetic position", "staff_type": "Synthetic type",
    }
    offering_columns = {
        field: f"Synthetic {field}" for field in schema_payloads[ModuleOffering]
    }
    team_columns = {
        "offering_id": "Synthetic offering", "staff_id": "Synthetic ID", "role": "Synthetic role",
    }
    config = ImportConfig(tables=[
        TableMapping(kind="staff", workbook=path, worksheet="Synthetic staff", columns=staff_columns),
        TableMapping(
            kind="preferences", workbook=path, worksheet="Synthetic preferences",
            columns={"staff_id": "Synthetic ID"},
            preference_fields={"source_note": {"column": "Synthetic preference", "value_type": "text"}},
        ),
        TableMapping(kind="offerings", workbook=path, worksheet="Synthetic offerings", columns=offering_columns),
        TableMapping(kind="teaching_teams", workbook=path, worksheet="Synthetic teams", columns=team_columns),
    ])
    staff = schema_payloads[StaffMember]
    offering = schema_payloads[ModuleOffering]
    records = [
        [staff, {**staff, "staff_id": "LEC090"}, {
            **staff, "staff_id": "INS007", "staff_name": "Example Instructor", "staff_type": "INSTRUCTOR",
        }],
        [["LEC001", "Private preference A"], ["LEC090", "Private preference B"]],
        [offering, {**offering, "offering_id": "OFFER002", "module_name": "Example Module B"}],
        [
            {"offering_id": "OFFER001", "staff_id": "LEC001", "role": "LIC"},
            {"offering_id": "OFFER001", "staff_id": "LEC090", "role": "CO_LIC"},
            {"offering_id": "OFFER002", "staff_id": "INS007", "role": "TEAM_MEMBER"},
        ],
    ]
    workbook = Workbook()
    workbook.remove(workbook.active)
    for table, table_records in zip(config.tables, records, strict=True):
        sheet = workbook.create_sheet(table.worksheet)
        sheet.append(list(table.columns.values()) + [rule.column for rule in table.preference_fields.values()])
        for record in table_records:
            sheet.append([record[field] for field in table.columns] if isinstance(record, dict) else record)
    workbook.save(path)
    workbook.close()
    return {"path": path, "config": config}


def set_cell(dataset, worksheet, coordinate, value):
    workbook = load_workbook(dataset["path"])
    workbook[worksheet][coordinate] = value
    workbook.save(dataset["path"])
    workbook.close()


def errors(result, code):
    return [issue for issue in result.issues if issue.severity == "error" and issue.code == code]


def test_valid_import_preserves_ids_conflicting_module_codes_and_source_file(excel_dataset):
    before = excel_dataset["path"].read_bytes()
    result = import_excel(excel_dataset["config"])
    assert result.valid
    assert [staff.staff_id for staff in result.academic_staff] == ["LEC001", "LEC090"]
    assert [staff.staff_id for staff in result.instructors] == ["INS007"]
    assert len(result.lecturer_preferences) == 2
    assert result.lecturer_preferences[0].values == {"source_note": "Private preference A"}
    assert [offering.module_code for offering in result.module_offerings] == ["IT2100", "IT2100"]
    assert [offering.offering_id for offering in result.module_offerings] == ["OFFER001", "OFFER002"]
    assert result.teaching_team_members[2].role is TeachingRole.TEAM_MEMBER
    assert len(result.teaching_team_members) == 3
    assert excel_dataset["path"].read_bytes() == before
    assert result.summary()["rows_read"] == {"staff": 3, "preferences": 2, "offerings": 2, "teaching_teams": 3}
    assert len([issue for issue in result.issues if issue.code == "shared_module_code"]) == 2
    assert not any(issue.code == "missing_lecturer_preference" for issue in result.issues)


def test_missing_column_reports_exact_location(excel_dataset):
    set_cell(excel_dataset, "Synthetic staff", "B1", "Unexpected header")
    result = import_excel(excel_dataset["config"])
    issue = errors(result, "missing_column")[0]
    assert (issue.workbook, issue.worksheet, issue.row, issue.column) == (
        "synthetic.xlsx", "Synthetic staff", 1, "Synthetic name",
    )
    assert result.staff_members == []


def test_duplicate_mapped_header_is_rejected(excel_dataset):
    set_cell(excel_dataset, "Synthetic staff", "E1", "Synthetic ID")
    assert errors(import_excel(excel_dataset["config"]), "duplicate_column")


@pytest.mark.parametrize("value", [None, " ", 123])
def test_staff_missing_or_invalid_text_reports_row_and_column(value, excel_dataset):
    set_cell(excel_dataset, "Synthetic staff", "B2", value)
    result = import_excel(excel_dataset["config"])
    issue = errors(result, "invalid_record")[0]
    assert issue.row == 2 and issue.column == "Synthetic name"
    assert "Example Lecturer" not in json.dumps(result.summary())
    assert errors(result, "unknown_staff")


@pytest.mark.parametrize(("worksheet", "coordinate"), [
    ("Synthetic teams", "B2"), ("Synthetic preferences", "A2"),
])
def test_unknown_staff_references(worksheet, coordinate, excel_dataset):
    set_cell(excel_dataset, worksheet, coordinate, "UNKNOWN-ID")
    result = import_excel(excel_dataset["config"])
    issue = errors(result, "unknown_staff")[0]
    assert issue.worksheet == worksheet and issue.row == 2 and issue.column == "Synthetic ID"


def test_unknown_offering_reference(excel_dataset):
    set_cell(excel_dataset, "Synthetic teams", "A2", "UNKNOWN-OFFERING")
    assert errors(import_excel(excel_dataset["config"]), "unknown_offering")


@pytest.mark.parametrize("role", ["LIC", "CO_LIC"])
def test_instructors_cannot_hold_academic_roles(role, excel_dataset):
    set_cell(excel_dataset, "Synthetic teams", "C4", role)
    result = import_excel(excel_dataset["config"])
    issue = errors(result, "invalid_role_assignment")[0]
    assert issue.row == 4 and issue.column == "Synthetic role"
    assert len(result.teaching_team_members) == 2


def test_instructor_preference_record_is_rejected(excel_dataset):
    set_cell(excel_dataset, "Synthetic preferences", "A2", "INS007")
    result = import_excel(excel_dataset["config"])
    assert errors(result, "instructor_preference")
    missing = [issue for issue in result.issues if issue.code == "missing_lecturer_preference"]
    assert len(missing) == 1 and missing[0].row == 2


def test_missing_academic_preference_is_a_warning_without_fabrication(excel_dataset):
    workbook = load_workbook(excel_dataset["path"])
    workbook["Synthetic preferences"].delete_rows(3)
    workbook.save(excel_dataset["path"])
    workbook.close()
    result = import_excel(excel_dataset["config"])
    assert result.valid and len(result.lecturer_preferences) == 1
    missing = [issue for issue in result.issues if issue.code == "missing_lecturer_preference"]
    assert len(missing) == 1 and missing[0].row == 3


@pytest.mark.parametrize("worksheet", [
    "Synthetic staff", "Synthetic preferences", "Synthetic offerings", "Synthetic teams",
])
def test_duplicate_identities_exclude_all_occurrences(worksheet, excel_dataset):
    workbook = load_workbook(excel_dataset["path"])
    sheet = workbook[worksheet]
    sheet.append([cell.value for cell in sheet[2]])
    workbook.save(excel_dataset["path"])
    workbook.close()
    result = import_excel(excel_dataset["config"])
    issues = errors(result, "duplicate_record")
    assert len(issues) == 2 and issues[0].row == 2
    assert not result.valid


def test_invalid_duplicate_staff_does_not_leave_an_ambiguous_eligible_record(excel_dataset):
    workbook = load_workbook(excel_dataset["path"])
    workbook["Synthetic staff"].append(["LEC001", None, "Instructor", "INSTRUCTOR"])
    workbook.save(excel_dataset["path"])
    workbook.close()
    result = import_excel(excel_dataset["config"])
    assert errors(result, "duplicate_record")
    assert "LEC001" not in {staff.staff_id for staff in result.staff_members}
    assert errors(result, "unknown_staff")


@pytest.mark.parametrize(("worksheet", "coordinate", "value"), [
    ("Synthetic staff", "D2", "UNKNOWN"),
    ("Synthetic offerings", "D2", 5),
    ("Synthetic offerings", "E2", 3),
    ("Synthetic offerings", "H2", "Evening"),
    ("Synthetic teams", "C2", "Assistant"),
    ("Synthetic staff", "A2", 1),
])
def test_existing_schema_validation_is_used(worksheet, coordinate, value, excel_dataset):
    set_cell(excel_dataset, worksheet, coordinate, value)
    assert errors(import_excel(excel_dataset["config"]), "invalid_record")


def test_explicit_enum_translations_do_not_infer_staff_type_or_alter_ids(excel_dataset):
    set_cell(excel_dataset, "Synthetic staff", "D2", "Source Academic")
    set_cell(excel_dataset, "Synthetic teams", "C2", "Source LIC")
    set_cell(excel_dataset, "Synthetic offerings", "H2", "Source Weekday")
    config = excel_dataset["config"]
    config.tables[0].value_maps = {"staff_type": {"Source Academic": "ACADEMIC"}}
    config.tables[2].value_maps = {"stream": {"Source Weekday": "WEEKDAY"}}
    config.tables[3].value_maps = {"role": {"Source LIC": "LIC"}}
    result = import_excel(config)
    assert result.valid and result.staff_members[0].staff_id == "LEC001"
    assert result.teaching_team_members[0].role is TeachingRole.LIC


@pytest.mark.parametrize(("value_type", "value", "valid"), [
    ("text", "Source text", True), ("text", 42, False),
    ("integer", 42, True), ("integer", True, False), ("integer", "42", False),
    ("number", 2.5, True), ("number", True, False),
    ("boolean", True, True), ("boolean", "yes", False),
])
def test_preference_types_are_explicit(value_type, value, valid, excel_dataset):
    rule = excel_dataset["config"].tables[1].preference_fields["source_note"]
    rule.value_type = value_type
    set_cell(excel_dataset, "Synthetic preferences", "B2", value)
    set_cell(excel_dataset, "Synthetic preferences", "B3", value)
    result = import_excel(excel_dataset["config"])
    assert result.valid is valid
    assert bool(errors(result, "invalid_preference_type")) is not valid


def test_missing_required_preference_is_reported(excel_dataset):
    set_cell(excel_dataset, "Synthetic preferences", "B2", None)
    issue = errors(import_excel(excel_dataset["config"]), "missing_value")[0]
    assert issue.row == 2 and issue.column == "Synthetic preference"


def test_optional_preference_is_omitted_without_a_default(excel_dataset):
    table = excel_dataset["config"].tables[1]
    table.preference_fields["optional_note"] = PreferenceField(
        column="Synthetic optional", value_type="text", required=False,
    )
    # Revalidate the configuration after adding a field to the mutable mapping.
    excel_dataset["config"] = ImportConfig.model_validate(excel_dataset["config"].model_dump())
    set_cell(excel_dataset, "Synthetic preferences", "C1", "Synthetic optional")
    result = import_excel(excel_dataset["config"])
    assert result.valid
    assert "optional_note" not in result.lecturer_preferences[0].values


@pytest.mark.parametrize("value", ["=1+1", "#DIV/0!"])
def test_formula_and_excel_error_cells_are_rejected(value, excel_dataset):
    set_cell(excel_dataset, "Synthetic preferences", "B2", value)
    issue = errors(import_excel(excel_dataset["config"]), "unsupported_cell")[0]
    assert issue.row == 2 and issue.column == "Synthetic preference"


def test_unmapped_columns_and_blank_rows_are_handled_explicitly(excel_dataset):
    set_cell(excel_dataset, "Synthetic staff", "E1", "Synthetic unused")
    set_cell(excel_dataset, "Synthetic staff", "E2", "Secret ignored content")
    workbook = load_workbook(excel_dataset["path"])
    workbook["Synthetic staff"].append([None, None, None, None])
    workbook.save(excel_dataset["path"])
    workbook.close()
    result = import_excel(excel_dataset["config"])
    assert result.valid and result.rows_read["staff"] == 3
    assert any(issue.code == "unmapped_columns" for issue in result.issues)
    assert "Secret ignored content" not in json.dumps(result.summary())


def test_configurable_header_row(excel_dataset):
    workbook = load_workbook(excel_dataset["path"])
    workbook["Synthetic staff"].insert_rows(1)
    workbook["Synthetic staff"]["A1"] = "Synthetic title"
    workbook.save(excel_dataset["path"])
    workbook.close()
    excel_dataset["config"].tables[0].header_row = 2
    assert import_excel(excel_dataset["config"]).valid


def test_missing_sheet_missing_file_corrupt_file_and_unsupported_extension(excel_dataset, tmp_path):
    table = deepcopy(excel_dataset["config"].tables[0])
    table.worksheet = "Absent"
    assert errors(import_excel(ImportConfig(tables=[table])), "missing_worksheet")
    table.workbook = tmp_path / "absent.xlsx"
    assert errors(import_excel(ImportConfig(tables=[table])), "unreadable_workbook")
    table.workbook.write_text("Not an Excel workbook", encoding="utf-8")
    assert errors(import_excel(ImportConfig(tables=[table])), "unreadable_workbook")
    table.workbook = tmp_path / "unsupported.xls"
    assert errors(import_excel(ImportConfig(tables=[table])), "unsupported_workbook")


@pytest.mark.parametrize("mutation", [
    "missing_mapping", "unknown_mapping", "duplicate_mapping", "id_translation", "bad_enum_destination",
    "preference_without_fields", "preference_unknown_column", "no_tables",
])
def test_invalid_mapping_configuration_is_rejected(mutation, excel_dataset):
    payload = excel_dataset["config"].model_dump(mode="json")
    table = payload["tables"][0]
    if mutation == "missing_mapping":
        del table["columns"]["staff_id"]
    elif mutation == "unknown_mapping":
        table["columns"]["unknown"] = "Synthetic unknown"
    elif mutation == "duplicate_mapping":
        table["columns"]["staff_name"] = table["columns"]["staff_id"]
    elif mutation == "id_translation":
        table["value_maps"] = {"staff_id": {"LEC001": "NEW001"}}
    elif mutation == "bad_enum_destination":
        table["value_maps"] = {"staff_type": {"Source": "INVALID"}}
    elif mutation == "preference_without_fields":
        payload["tables"][1]["preference_fields"] = {}
    elif mutation == "preference_unknown_column":
        payload["tables"][1]["columns"]["inferred_availability"] = "Synthetic unknown"
    else:
        payload["tables"] = []
    with pytest.raises(ValidationError):
        ImportConfig.model_validate(payload)


def write_config(dataset, tmp_path):
    payload = dataset["config"].model_dump(mode="json")
    for table in payload["tables"]:
        table["workbook"] = dataset["path"].name
    path = tmp_path / "mapping.json"
    path.write_text(json.dumps(payload), encoding="utf-8")
    return path


def test_cli_relative_paths_counts_privacy_and_exit_status(excel_dataset, tmp_path, capsys):
    config_path = write_config(excel_dataset, tmp_path)
    assert main(["--config", str(config_path)]) == 0
    output = capsys.readouterr().out
    summary = json.loads(output)
    assert summary["valid"] and summary["data_status"] == "development_unverified"
    assert summary["imported_counts"] == {
        "staff_members": 3, "academic_staff": 2, "instructors": 1,
        "lecturer_preferences": 2, "module_offerings": 2, "teaching_team_members": 3,
    }
    for sensitive in ("LEC001", "LEC090", "INS007", "Example Lecturer", "Private preference A"):
        assert sensitive not in output
    set_cell(excel_dataset, "Synthetic teams", "B2", "UNKNOWN")
    assert main(["--config", str(config_path)]) == 1
    assert json.loads(capsys.readouterr().out)["errors"]


def test_module_command_runs_in_a_real_subprocess(excel_dataset, tmp_path):
    config_path = write_config(excel_dataset, tmp_path)
    completed = subprocess.run(
        [sys.executable, "-m", "schedence_teaching.excel_import.cli", "--config", str(config_path)],
        capture_output=True, text=True, check=False,
    )
    assert completed.returncode == 0, completed.stderr
    assert json.loads(completed.stdout)["imported_counts"]["module_offerings"] == 2


@pytest.mark.parametrize("content", ["not json", '{"tables": []}', '{"tables": [{"kind": "staff"}]}'])
def test_invalid_cli_configuration_has_a_clear_failure_status(content, tmp_path, capsys):
    path = tmp_path / "mapping.json"
    path.write_text(content, encoding="utf-8")
    assert main(["--config", str(path)]) == 2
    assert json.loads(capsys.readouterr().out)["errors"][0]["code"] == "invalid_config"


def test_missing_cli_configuration_is_reported(tmp_path, capsys):
    assert main(["--config", str(tmp_path / "absent.json")]) == 2
    assert json.loads(capsys.readouterr().out)["errors"][0]["code"] == "unreadable_config"
