"""Sanitized fixtures use inspected headers, never real staff names or preferences."""

from __future__ import annotations

import json

import pytest
from openpyxl import Workbook, load_workbook

from schedence_teaching.excel_import import WideImportConfig, import_wide_excel
from schedence_teaching.excel_import.cli import main

STAFF_HEADERS = ["Staff_ID", "Staff_Name", "Position", "Data_Status", "Department", "Department_Source"]
PREFERENCE_HEADERS = [
    "Staff_ID", "Staff_Name", "Department", "Position", "Expected_Working_Days", "Weekend_Work_Pattern",
    "Preferred_Free_Day", "Can_Work_Both_Weekend_Days", "Monday_Morning", "Monday_Afternoon",
    "Tuesday_Morning", "Tuesday_Afternoon", "Wednesday_Morning", "Wednesday_Afternoon", "Thursday_Morning",
    "Thursday_Afternoon", "Friday_Morning", "Friday_Afternoon", "Saturday_Morning", "Saturday_Afternoon",
    "Sunday_Morning", "Sunday_Afternoon", "Special_Notes",
]
MODULE_HEADERS = [
    "Module_Code", "Module_Name", "Year", "Semester", "Specialization", "LIC", "Co_LIC",
    "Teaching_Team_Member_1", "Teaching_Team_Member_2", "Teaching_Team_Member_3", "Teaching_Team_Member_4",
    "Teaching_Team_Member_5", "Module_Degree_Type", "Module_Credits", "Exam_Type", "Exam_Duration_Minutes",
    "Extra_Reading_Time", "Difficulty_Level", "Has_Tutorial", "Has_Workshop", "Has_Lab", "Lecture_Hours_Week",
    "Tutorial_Hours_Week", "Workshop_Hours_Week", "Lab_Hours_Week", "Total_Student_Contact_Hours",
    "LIC_Staff_ID", "Co_LIC_Staff_ID", "Team_Staff_ID_1", "Team_Staff_ID_2", "Team_Staff_ID_3",
    "Team_Staff_ID_4", "Team_Staff_ID_5", "Lecturer Selection Options (90 academic + 20 DEMO instructors)",
]


@pytest.fixture
def wide_dataset(tmp_path):
    staff = []
    for index, identity in enumerate(("LEC001", "LEC090", "LEC003", "LEC004", "LEC005", "INS001", "INS002")):
        staff.append({
            "Staff_ID": identity, "Staff_Name": f"Synthetic Person {index}",
            "Position": "Lecturer" if index < 5 else "Instructor",
            "Data_Status": "Academic source" if index < 5 else "Demo instructor source",
            "Department": "Synthetic Department", "Department_Source": "Synthetic development provenance",
        })
    paths = {}
    for key, sheet_name, headers, header_row in (
        ("staff", "Staff_Master", STAFF_HEADERS, 1),
        ("preferences", "Lecturer_Preferences", PREFERENCE_HEADERS, 4),
        ("modules", "Modules_183", MODULE_HEADERS, 1),
    ):
        path = tmp_path / f"synthetic-{key}.xlsx"
        paths[key] = path
        book = Workbook()
        sheet = book.active
        sheet.title = sheet_name
        for index in range(header_row - 1):
            sheet.append([f"Synthetic template heading {index}"])
        sheet.append(headers)
        if key == "staff":
            records = staff
        elif key == "preferences":
            records = [{
                **record,
                **{name: "Available" for name in PREFERENCE_HEADERS[8:22]},
                "Expected_Working_Days": 5, "Weekend_Work_Pattern": "Saturday Only",
                "Preferred_Free_Day": "Sunday", "Can_Work_Both_Weekend_Days": "No",
                "Special_Notes": "Private synthetic preference",
            } for record in staff[:5]]
        else:
            records = []
            for index in range(2):
                record = {
                    "Module_Code": "IT2100", "Module_Name": f"Synthetic Module {index}", "Year": 2,
                    "Semester": f"Semester {index + 1}", "Specialization": "Synthetic Specialization",
                    "Module_Degree_Type": "Local", "Module_Credits": 3, "Exam_Type": "written",
                    "Exam_Duration_Minutes": 120, "Extra_Reading_Time": "No", "Difficulty_Level": 2,
                    "Has_Tutorial": "Yes", "Has_Workshop": "No", "Has_Lab": "Yes",
                    "Lecture_Hours_Week": 2, "Tutorial_Hours_Week": 1, "Workshop_Hours_Week": 0,
                    "Lab_Hours_Week": 2, "Total_Student_Contact_Hours": f"=SUM(V{index+2}:Y{index+2})",
                }
                pairs = [("LIC_Staff_ID", "LIC"), ("Co_LIC_Staff_ID", "Co_LIC")] + [
                    (f"Team_Staff_ID_{number}", f"Teaching_Team_Member_{number}") for number in range(1, 6)
                ]
                for person, (id_column, name_column) in zip(staff, pairs, strict=True):
                    record[id_column] = person["Staff_ID"]
                    record[name_column] = f'{person["Position"]} | {person["Staff_Name"]}'
                records.append(record)
        for record in records:
            sheet.append([record.get(header) for header in headers])
        book.save(path)
        book.close()
    config = WideImportConfig(
        layout="wide-teaching-v1",
        staff={"workbook": paths["staff"], "worksheet": "Staff_Master", "header_row": 1},
        preferences={"workbook": paths["preferences"], "worksheet": "Lecturer_Preferences", "header_row": 4},
        modules={"workbook": paths["modules"], "worksheet": "Modules_183", "header_row": 1},
        staff_type_values={"Academic source": "ACADEMIC", "Demo instructor source": "INSTRUCTOR"},
        semester_values={"Semester 1": 1, "Semester 2": 2},
    )
    return {"config": config, "paths": paths, "staff": staff}


def change(dataset, key, header, value, row=None):
    path = dataset["paths"][key]
    header_row = 4 if key == "preferences" else 1
    book = load_workbook(path)
    sheet = book.active
    headers = [cell.value for cell in sheet[header_row]]
    sheet.cell(row or header_row + 1, headers.index(header) + 1, value)
    if value is None:
        sheet.cell(row or header_row + 1, headers.index(header) + 1).value = None
    book.save(path)
    book.close()


def codes(result):
    return {issue.code for issue in result.issues}


def test_wide_normalization_preserves_every_assignment_without_invented_ids(wide_dataset):
    before = {key: path.read_bytes() for key, path in wide_dataset["paths"].items()}
    result = import_wide_excel(wide_dataset["config"])
    assert result.valid
    assert len(result.staff_members) == 7 and len(result.academic_staff) == 5 and len(result.instructors) == 2
    assert len(result.lecturer_preferences) == 5
    assert len(result.source_modules) == 2 and len(result.source_assignments) == 14
    assert all(item.valid_reference for item in result.source_assignments)
    assert [item.contract_fields["semester"] for item in result.source_modules] == [1, 2]
    assert all(item.contract_fields["module_code"] == "IT2100" for item in result.source_modules)
    assert result.module_offerings == result.teaching_team_members == []
    assert result.source_modules[0].location.diagnostic_key.startswith("source-row:")
    assert "offering_id" not in result.source_modules[0].contract_fields
    for module in result.source_modules:
        assert module.source_values["Total_Student_Contact_Hours"].startswith("=SUM(")
    for assignment in result.source_assignments:
        source = next(item for item in result.source_modules if item.location == assignment.location)
        assert assignment.source_staff_id == source.source_values[assignment.id_column]
        assert assignment.source_display_name == source.source_values[assignment.name_column]
    assert result.lecturer_preferences[0].values["Can_Work_Both_Weekend_Days"] == "No"
    assert not result.summary()["ready_for_scheduling"]
    assert {key: path.read_bytes() for key, path in wide_dataset["paths"].items()} == before


@pytest.mark.parametrize(("key", "header", "value", "expected"), [
    ("modules", "Year", 5, "invalid_record"),
    ("modules", "Semester", "Semester 3", "invalid_record"),
    ("modules", "LIC_Staff_ID", "UNKNOWN", "unknown_staff"),
    ("modules", "LIC", "Different source display", "assignment_name_mismatch"),
    ("modules", "LIC_Staff_ID", "INS001", "invalid_role_assignment"),
    ("modules", "Co_LIC_Staff_ID", "INS001", "invalid_role_assignment"),
    ("modules", "Team_Staff_ID_1", "LEC001", "repeated_team_staff"),
    ("preferences", "Staff_ID", "UNKNOWN", "unknown_staff"),
    ("preferences", "Staff_Name", "Other synthetic name", "preference_staff_mismatch"),
    ("preferences", "Department", "Other department", "preference_staff_mismatch"),
    ("preferences", "Expected_Working_Days", "5", "invalid_preference_type"),
    ("staff", "Data_Status", "Undeclared classification", "invalid_record"),
    ("modules", "LIC_Staff_ID", "=1+1", "unsupported_cell"),
])
def test_source_errors_have_locations_without_rewriting_assignments(key, header, value, expected, wide_dataset):
    change(wide_dataset, key, header, value)
    result = import_wide_excel(wide_dataset["config"])
    assert not result.valid and expected in codes(result)
    assert len(result.source_assignments) == 14
    issue = next(issue for issue in result.issues if issue.code == expected)
    assert issue.workbook and issue.worksheet and issue.row is not None and issue.column


@pytest.mark.parametrize("id_column,name_column", [("LIC_Staff_ID", "LIC"), ("Co_LIC_Staff_ID", "Co_LIC")])
def test_missing_lead_is_reported_without_substitution(id_column, name_column, wide_dataset):
    change(wide_dataset, "modules", id_column, None)
    change(wide_dataset, "modules", name_column, None)
    result = import_wide_excel(wide_dataset["config"])
    assert "missing_lead_assignment" in codes(result)
    assert len(result.source_assignments) == 13


def test_same_id_different_name_and_duplicate_ids_are_reported(wide_dataset):
    change(wide_dataset, "staff", "Staff_ID", "LEC001", row=3)
    result = import_wide_excel(wide_dataset["config"])
    assert {"staff_id_name_conflict", "duplicate_record"}.issubset(codes(result))
    assert len(result.source_staff) == 7 and len(result.source_assignments) == 14


def test_same_name_different_ids_are_flagged_without_merging_people(wide_dataset):
    change(wide_dataset, "staff", "Staff_Name", "Synthetic Person 0", row=3)
    result = import_wide_excel(wide_dataset["config"])
    assert "same_name_multiple_ids" in codes(result)
    assert len(result.staff_members) == 7


def test_duplicate_module_rows_are_retained_for_review(wide_dataset):
    book = load_workbook(wide_dataset["paths"]["modules"])
    book.active.append([cell.value for cell in book.active[2]])
    book.save(wide_dataset["paths"]["modules"])
    book.close()
    result = import_wide_excel(wide_dataset["config"])
    assert "duplicate_module_record" in codes(result)
    assert len(result.source_modules) == 3 and len(result.source_assignments) == 21


def test_identity_ambiguity_is_reported_without_dropping_different_teams(wide_dataset):
    change(wide_dataset, "modules", "Module_Name", "Synthetic Module 0", row=3)
    change(wide_dataset, "modules", "Semester", "Semester 1", row=3)
    result = import_wide_excel(wide_dataset["config"])
    assert "ambiguous_module_identity" in codes(result)
    assert len(result.source_assignments) == 14


def test_missing_preferences_are_never_generated(wide_dataset):
    book = load_workbook(wide_dataset["paths"]["preferences"])
    book.active.delete_rows(5)
    book.save(wide_dataset["paths"]["preferences"])
    book.close()
    result = import_wide_excel(wide_dataset["config"])
    assert result.valid and len(result.lecturer_preferences) == 4
    assert "missing_lecturer_preference" in codes(result)
    assert len([issue for issue in result.issues if issue.code == "missing_lecturer_preference"]) == 1


def test_no_department_or_stream_is_inferred(wide_dataset):
    change(wide_dataset, "staff", "Department", None)
    result = import_wide_excel(wide_dataset["config"])
    assert result.staff_members[0].department is None
    assert all("stream" not in item.contract_fields for item in result.source_modules)
    assert any(blocker.get("field") == "stream" for blocker in result.summary()["contract_blockers"])


def test_actual_offering_identifiers_and_streams_enable_contract_conversion(wide_dataset):
    book = load_workbook(wide_dataset["paths"]["modules"])
    sheet = book.active
    sheet.cell(1, 35, "Authoritative_Offering_ID")
    sheet.cell(1, 36, "Verified_Stream")
    for row in (2, 3):
        sheet.cell(row, 35, f"AUTHORITATIVE-{row}")
        sheet.cell(row, 36, "WEEKDAY")
    book.save(wide_dataset["paths"]["modules"])
    book.close()
    config = wide_dataset["config"]
    config.offering_id_column = "Authoritative_Offering_ID"
    config.stream_column = "Verified_Stream"
    result = import_wide_excel(config)
    assert result.valid and len(result.module_offerings) == 2 and len(result.teaching_team_members) == 14
    assert {row.offering_id for row in result.module_offerings} == {"AUTHORITATIVE-2", "AUTHORITATIVE-3"}
    assert "contract_conversion_blocked" not in codes(result)


def test_cli_summary_is_private_and_distinguishes_staging_from_contracts(wide_dataset, tmp_path, capsys):
    path = tmp_path / "mapping.json"
    path.write_text(wide_dataset["config"].model_dump_json(), encoding="utf-8")
    assert main(["--config", str(path)]) == 0
    output = capsys.readouterr().out
    assert "Synthetic Person" not in output and "Private synthetic preference" not in output
    assert "LEC001" not in output and "INS001" not in output
    summary = json.loads(output)
    assert summary["staged_counts"]["teaching_team_assignments"] == 14
    assert summary["imported_counts"]["teaching_team_members"] == 0
    assert not summary["ready_for_scheduling"]


@pytest.mark.parametrize("problem", ["missing_header", "duplicate_header", "missing_sheet", "missing_file"])
def test_wide_source_read_failures_are_explicit(problem, wide_dataset):
    config = wide_dataset["config"]
    if problem in {"missing_header", "duplicate_header"}:
        book = load_workbook(wide_dataset["paths"]["modules"])
        book.active.cell(1, 1, "Absent header" if problem == "missing_header" else "Module_Name")
        book.save(wide_dataset["paths"]["modules"])
        book.close()
    elif problem == "missing_sheet":
        config.modules.worksheet = "Absent sheet"
    else:
        config.modules.workbook = config.modules.workbook.parent / "absent.xlsx"
    result = import_wide_excel(config)
    assert not result.valid and result.source_modules == []
    assert codes(result).intersection({"missing_column", "duplicate_column", "missing_worksheet", "unreadable_workbook"})


def test_staff_type_comes_from_declared_status_not_id_prefix_or_position(wide_dataset):
    change(wide_dataset, "staff", "Data_Status", "Demo instructor source")
    result = import_wide_excel(wide_dataset["config"])
    assert result.staff_members[0].staff_id == "LEC001"
    assert result.staff_members[0].position == "Lecturer"
    assert result.staff_members[0].staff_type.value == "INSTRUCTOR"
    assert {"invalid_role_assignment", "instructor_preference"}.issubset(codes(result))
