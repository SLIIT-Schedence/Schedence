"""Editable-layout and preview tests use synthetic workbooks/reference records."""

import csv
from dataclasses import replace
from hashlib import sha256
from itertools import product
import json

from openpyxl import Workbook, load_workbook
import pytest

from schedence_teaching.excel_import import VenueImportResult
from schedence_teaching.excel_import.staging import SourceLocation, SourceModuleRecord, SourceRecord, SourceTeamAssignment, WideImportResult
from schedence_teaching.integration import SchedulingSubmission, assess_readiness
from schedence_teaching.pilot.demo import DEMO_ASSUMPTIONS, DEMO_INSTITUTION, build_demo_submission
from schedence_teaching.pilot.preview import PilotMode, export_validated_preview, run_pilot_preview
from schedence_teaching.pilot.students import (
    CohortCount, GroupSizingPolicy, balanced_sizes, generate_groups, import_student_counts,
)
from schedence_teaching.preparation import build_group_hierarchy
from schedence_teaching.schemas import StaffMember, TeachingRole, Venue
from schedence_teaching.solver import TimetableValidationError


# Fixture snapshot of the user-confirmed product inputs, never engine constants.
COUNTS = ((720, 120), (720, 120), (600, 240), (600, 240), (480, 360), (240, 600), (240, 600), (240, 600))


@pytest.fixture
def workbook_path(tmp_path):
    file = tmp_path / "synthetic-student-input.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "STUDENT_GROUPS"
    for address, value in {"A4": "Main target", "B4": 120, "C4": "Lab cap", "D4": 60,
                           "G4": "Large hall", "H4": 240}.items():
        sheet[address] = value
    for index, header in enumerate(("Cohort ID", "Study year", "Semester", "Stream", "Specialization", "Students",
        "Main groups", "Lab groups", "Readiness", "Data source"), 1):
        sheet.cell(9, index, header)
    index = 0
    for (year, semester), totals in zip(product(range(1, 5), (1, 2)), COUNTS):
        for stream, count in zip(("WEEKDAY", "WEEKEND"), totals):
            values = (f"SYN-Y{year}-S{semester}-{stream}", year, semester, stream, None, count,
                      "=1+1", "=2+2", "=1", "Synthetic development counts")
            for column, value in enumerate(values, 1):
                sheet.cell(10+index, column, value)
            index += 1
    for column in range(1, 11):
        sheet.cell(33, column, "=1+1")  # Generated output is never imported as cohort input.
    book.save(file)
    book.close()
    return file


def edit(file, address, value):
    book = load_workbook(file)
    book["STUDENT_GROUPS"][address] = value
    book.save(file)
    book.close()


def cohort(total=720, **kwargs):
    return CohortCount(cohort_id="SYN-COHORT", academic_year=1, semester=1, stream="WEEKDAY", student_count=total,
        source=SourceLocation("synthetic.xlsx", "STUDENT_GROUPS", 10), **kwargs)


def test_import_all_counts_scopes_and_source_unchanged(workbook_path):
    before = sha256(workbook_path.read_bytes()).hexdigest()
    imported = import_student_counts(workbook_path)
    assert imported.valid, imported.summary()
    assert len(imported.cohorts) == 16
    assert sum(row.student_count for row in imported.cohorts) == 6720
    assert {(row.academic_year, row.semester, row.stream.value) for row in imported.cohorts} == set(product(range(1, 5), (1, 2), ("WEEKDAY", "WEEKEND")))
    assert all(row.specialization is None for row in imported.cohorts)
    assert not any(row.code == "unsupported_cell" for row in imported.issues)
    assert imported.source_sha256 == before == sha256(workbook_path.read_bytes()).hexdigest()


def test_all_planned_counts_and_parent_populations_not_double_counted(workbook_path):
    imported = import_student_counts(workbook_path)
    plan = generate_groups(imported.cohorts, imported.sizing)
    summary = plan.summary()
    assert (summary["total_students"], summary["main_group_count"], summary["lab_subgroup_count"]) == (6720, 56, 112)
    assert summary["group_record_count"] == 168
    assert summary["main_student_sum"] == summary["lab_student_sum"] == 6720
    assert len({row.group_id for row in plan.groups}) == 168
    assert plan.specializations == {}
    assert not summary["population_disjointness_verified"]
    hierarchy = build_group_hierarchy(plan.groups)
    for parent in (row for row in plan.groups if row.parent_group_id is None):
        children = [row for row in plan.groups if row.parent_group_id == parent.group_id]
        assert sum(row.student_count for row in children) == parent.student_count <= 120
        assert all(row.student_count <= 60 and hierarchy.conflicts(parent.group_id, row.group_id) for row in children)


def test_first_pilot_counts_and_deterministic_identity():
    first = generate_groups([cohort()], GroupSizingPolicy())
    second = generate_groups([cohort()], GroupSizingPolicy())
    assert first.groups == second.groups
    assert first.summary()["total_students"] == 720
    assert (first.summary()["main_group_count"], first.summary()["lab_subgroup_count"]) == (6, 12)


def test_uneven_260_respects_caps():
    plan = generate_groups([cohort(260)], GroupSizingPolicy())
    main = [row.student_count for row in plan.groups if row.parent_group_id is None]
    assert main == [87, 87, 86]
    assert all(row.student_count <= (120 if row.parent_group_id is None else 60) for row in plan.groups)
    assert plan.summary()["total_students"] == 260


@pytest.mark.parametrize("total,maximum,expected", [(1, 120, (1,)), (121, 120, (61, 60)), (120, 120, (120,)), (260, 120, (87, 87, 86))])
def test_balanced_sizes(total, maximum, expected):
    assert balanced_sizes(total, maximum) == expected


@pytest.mark.parametrize("total,maximum", [(0, 120), (-1, 120), (True, 120), (1, 0), (2.5, 120)])
def test_invalid_sizing_never_generates_groups(total, maximum):
    with pytest.raises(ValueError):
        balanced_sizes(total, maximum)


def test_configurable_caps_do_not_impose_two_subgroups():
    plan = generate_groups([cohort(200)], GroupSizingPolicy(main_group_maximum=200, lab_subgroup_maximum=60))
    assert plan.summary()["main_group_count"] == 1
    assert plan.summary()["lab_subgroup_count"] == 4
    assert [row.student_count for row in plan.groups if row.parent_group_id] == [50]*4


def test_technical_guard_rejects_without_truncating_or_allocating_huge_partitions():
    with pytest.raises(ValueError, match="limit"):
        generate_groups([cohort(10**9)], GroupSizingPolicy(main_group_maximum=10**9, lab_subgroup_maximum=1))


@pytest.mark.parametrize("address,value", [("F10", 0), ("F10", "720"), ("F10", "=720"), ("B10", 5), ("C10", 3),
    ("D10", "EVENING"), ("A10", None), ("B4", 0), ("D4", "=60"), ("H4", -1)])
def test_invalid_literal_workbook_cells(workbook_path, address, value):
    edit(workbook_path, address, value)
    imported = import_student_counts(workbook_path)
    assert not imported.valid
    assert any(row.severity == "error" and row.row is not None for row in imported.issues)


def test_confirmed_intake_mapping(workbook_path):
    edit(workbook_path, "C10", "Regular Intake")
    edit(workbook_path, "C12", "June Intake")
    imported = import_student_counts(workbook_path)
    assert imported.valid
    assert imported.cohorts[0].semester == 1
    assert imported.cohorts[2].semester == 2


def test_edited_product_count_does_not_require_fixed_6720(workbook_path):
    edit(workbook_path, "F10", 260)
    imported = import_student_counts(workbook_path)
    assert imported.valid
    assert sum(row.student_count for row in imported.cohorts) == 6260


def test_only_supplied_specialization_is_propagated():
    plan = generate_groups([cohort(specialization="SYN-SUPPLIED")], GroupSizingPolicy())
    assert set(plan.specializations.values()) == {"SYN-SUPPLIED"}
    assert len(plan.specializations) == len(plan.groups)


def reference_sources():
    academic = WideImportResult()
    for identity, position, kind in (("SYN-A", "Lecturer", "ACADEMIC"), ("SYN-I", "Instructor", "INSTRUCTOR"),
                                   ("SYN-B", "Assistant Lecturer", "ACADEMIC"), ("SYN-J", "Instructor", "INSTRUCTOR")):
        academic.staff_members.append(StaffMember(staff_id=identity, staff_name=f"Synthetic {identity}", position=position, staff_type=kind))
    for index, pool in enumerate((("SYN-A", "SYN-I"), ("SYN-B", "SYN-J")), 1):
        location = SourceLocation("synthetic-modules.xlsx", "Modules", index+1)
        academic.source_modules.append(SourceModuleRecord(location, {}, contract_fields=dict(module_code=f"SYN-M{index}",
            module_name="Synthetic <Module & Topic>", academic_year=1, semester=1, specialization="SYN-SOURCE", degree_type="Synthetic")))
        for member_index, identity in enumerate(pool):
            academic.source_assignments.append(SourceTeamAssignment(location, "ID", "Name", identity, "Synthetic",
                TeachingRole.LIC if member_index == 0 else TeachingRole.TEAM_MEMBER))
    venues = VenueImportResult()
    for index, (identity, capacity, kind) in enumerate((("SYN-H240", 240, "Lecture Hall"), ("SYN-H120", 120, "Lecture Hall"),
        ("SYN-L1", 60, "Computer Lab"), ("SYN-L2", 60, "Computer Lab")), 2):
        venues.venues.append(Venue(venue_id=identity, venue_name=identity, capacity=capacity, venue_type=kind))
        venues.source_venues.append(SourceRecord(SourceLocation("synthetic-venues.xlsx", "Venue_Master", index), {}))
    return academic, venues


@pytest.fixture(scope="module")
def solved_pilot(tmp_path_factory):
    plan = generate_groups([cohort()], GroupSizingPolicy())
    academic, venues = reference_sources()
    submission = build_demo_submission(plan, academic, venues, time_limit=10)
    path = tmp_path_factory.mktemp("synthetic-pilot") / "preview"
    run = run_pilot_preview(submission, plan, path, institution_id=DEMO_INSTITUTION, mode=PilotMode.DEMO,
        source_report=dict(synthetic_test_only=True), assumptions=DEMO_ASSUMPTIONS)
    assert run.label == "SYNTHETIC_DEMO", run.solver_status
    assessment = assess_readiness(submission, expected_institution_id=DEMO_INSTITUTION, actual_campus=False)
    return plan, submission, run, assessment.preparation


def test_native_demo_preserves_coverage_capacity_and_lunch(solved_pilot):
    plan, submission, run, prepared = solved_pilot
    assert run.solver_status in {"FEASIBLE", "OPTIMAL"}
    result = run.result
    coverage = [row.occurrence_id for delivery in result.schedule for row in delivery.covered_sessions]
    assert len(coverage) == len(set(coverage)) == 18
    assert result.student_lunch_report.assessment.compliant
    assert result.validation.expected_occurrences == result.validation.assigned_occurrences == 18
    capacities = {row.venue_id: row.capacity for row in prepared.solver_inputs["venue_input"].venues}
    for row in result.schedule:
        assert row.combined_student_count <= capacities[row.assigned_venue_id]
        assert not (row.start_minute < 810 and 750 < row.end_minute)
    assert plan.specializations == {}  # The real/source plan was never mutated.
    assert len(set(submission.inputs["combination_input"]["group_specializations"].values())) == 2


def test_html_csv_json_export_validated_deliveries_only(solved_pilot):
    _, _, run, _ = solved_pilot
    with open(run.paths["csv"], encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == len(run.result.schedule)
    assert {row["delivery_id"] for row in rows} == {row.delivery_id for row in run.result.schedule}
    assert {row["output_label"] for row in rows} == {"SYNTHETIC_DEMO"}
    html = open(run.paths["html"], encoding="utf-8").read()
    assert "SYNTHETIC_DEMO" in html and "12:30–13:30" in html
    assert "Monday" in html and "Wednesday" in html
    assert "&lt;Module &amp; Topic&gt;" in html
    assert "<Module & Topic>" not in html
    payload = json.load(open(run.paths["solver_json"], encoding="utf-8"))
    assert len(payload["solver"]["schedule"]) == len(rows)
    assert payload["solver"]["student_lunch"]["compliant"] is True


def test_corrupted_timetable_cannot_be_exported(solved_pilot, tmp_path):
    plan, _, run, preparation = solved_pilot
    bad = replace(run.result, schedule=run.result.schedule[:-1])
    output = tmp_path / "rejected"
    with pytest.raises(ValueError):
        export_validated_preview(output, preparation, bad, label="SYNTHETIC_DEMO", plan=plan, report={})
    assert not output.exists()


def test_synthetic_timetable_cannot_be_labelled_campus_validated(solved_pilot, tmp_path):
    plan, _, run, preparation = solved_pilot
    with pytest.raises(ValueError, match="label"):
        export_validated_preview(tmp_path / "wrong-label", preparation, run.result, label="CAMPUS_INPUT_VALIDATED", plan=plan, report={})


def test_real_mode_blocks_missing_inputs_and_never_calls_solver(tmp_path, monkeypatch):
    monkeypatch.setattr("schedence_teaching.integration.service.solve_timetable", lambda *a, **k: pytest.fail("Actual inputs are incomplete"))
    plan = generate_groups([cohort()], GroupSizingPolicy())
    output = run_pilot_preview(SchedulingSubmission(inputs={}, evidence=[]), plan, tmp_path / "blocked",
        institution_id=DEMO_INSTITUTION, mode=PilotMode.REAL_CAMPUS, source_report={})
    assert output.solver_status == "BLOCKED" and output.label is None
    assert set(output.paths) == {"readiness_json"}
    assert any(row.code == "student_plan_mapping" for row in output.readiness.diagnostics)


def test_demo_does_not_activate_real_campus_mode(solved_pilot, tmp_path):
    plan, submission, _, _ = solved_pilot
    output = run_pilot_preview(submission, plan, tmp_path / "real-rejected", institution_id=DEMO_INSTITUTION,
        mode=PilotMode.REAL_CAMPUS, source_report={})
    assert output.label is None and output.solver_status == "BLOCKED"


def test_exports_refuse_to_overwrite_existing_outputs(solved_pilot):
    plan, _, run, preparation = solved_pilot
    from pathlib import Path
    file = Path(run.paths["html"])
    before = file.read_bytes()
    with pytest.raises(FileExistsError):
        export_validated_preview(file.parent, preparation, run.result, label="SYNTHETIC_DEMO", plan=plan, report={})
    assert file.read_bytes() == before
